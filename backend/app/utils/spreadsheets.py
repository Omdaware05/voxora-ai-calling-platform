"""CSV / XLSX import-export helpers."""
from __future__ import annotations

import csv
import io
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook


def parse_csv(data: bytes) -> List[Dict[str, Any]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def parse_xlsx(data: bytes, sheet: int = 0) -> List[Dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
    out: List[Dict[str, Any]] = []
    for row in rows[1:]:
        if all(v is None or v == "" for v in row):
            continue
        out.append({h: ("" if v is None else v) for h, v in zip(headers, row)})
    return out


def to_csv(rows: List[Dict[str, Any]], fieldnames: List[str]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fieldnames)
    writer.writeheader()
    for row in rows:
        writer.writerow({k: row.get(k, "") for k in fieldnames})
    return buf.getvalue().encode("utf-8")


def questionnaire_template_xlsx() -> bytes:
    """Sample questionnaire import template as XLSX bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "questionnaire"
    ws.append(["title", "Customer Satisfaction Survey"])
    ws.append(["description", "Post-call satisfaction survey"])
    ws.append([])
    ws.append(["key", "label", "description", "type", "required", "order", "options", "expected_answer"])
    ws.append(["overall_rating", "Overall rating", "Rate the call 1-5", "rating", "TRUE", 1, "1,2,3,4,5", ""])
    ws.append(["issue_resolved", "Was your issue resolved?", "", "boolean", "TRUE", 2, "", "yes"])
    ws.append(["feedback", "Any additional feedback?", "", "text", "FALSE", 3, "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_questionnaire_workbook(data: bytes) -> Dict[str, Any]:
    """Parse the questionnaire template format into a questionnaire payload.

    Format (see questionnaire_template_xlsx):
      - optional metadata rows: ["title", <value>] / ["description", <value>]
      - a header row starting with "key","label",... ("type","required","order",
        "options","expected_answer",...)
      - one question per row after the header.
    """
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    raw_rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]

    title = "Imported Questionnaire"
    description: str | None = None
    questions: List[Dict[str, Any]] = []
    header: List[str] = []
    in_questions = False

    for row in raw_rows:
        cells = ["" if c is None else str(c).strip() for c in row]
        if not any(cells):
            continue
        first = cells[0].lower()
        if not in_questions and first in ("title", "description") and len(cells) > 1:
            if first == "title" and cells[1]:
                title = cells[1]
            elif first == "description" and cells[1]:
                description = cells[1]
            continue
        if not in_questions and first == "key" and any(c.lower() == "label" for c in cells):
            header = [c.lower() for c in cells]
            in_questions = True
            continue
        if not in_questions:
            continue
        record = {h: (cells[i] if i < len(cells) else "") for i, h in enumerate(header)}
        key = record.get("key", "").strip()
        if not key:
            continue
        required = record.get("required", "").upper() in ("TRUE", "YES", "1")
        options_raw = record.get("options", "").strip()
        try:
            order = int(record.get("order", "") or 0)
        except ValueError:
            order = 0
        qtype = record.get("type", "").strip().lower() or "text"
        questions.append(
            {
                "key": key,
                "label": record.get("label", "").strip() or key,
                "description": record.get("description", "").strip() or None,
                "type": qtype,
                "required": required,
                "order": order,
                "options": [o.strip() for o in options_raw.split(",") if o.strip()] or None,
                "validation": None,
                "conditional": None,
                "expected_answer": record.get("expected_answer", "").strip() or None,
            }
        )
    return {"title": title, "description": description, "questions": questions}
